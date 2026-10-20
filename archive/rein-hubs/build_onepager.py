"""Internal one-pager: REIN Ghana Hub x Academic City, for Laurence Kalinsky.

Facts are limited to the 20 Oct 2026 brief (source: REIN x Academic City
proposal, April 2026). Rebuild with:

    python3 build_onepager.py
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
LOGO = HERE / "rein-logo-header.png"
OUT = HERE / "2026-10-20-laurence-ghana-onepager.docx"

GREEN = RGBColor(0x33, 0x5D, 0x4F)
ORANGE = RGBColor(0xFF, 0x66, 0x33)
INK = RGBColor(0x22, 0x2A, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Stated cash figures only. Bars use these, with Year 1 as the full width.
YEAR1 = 103_000
YEAR2 = 75_000
YEAR3 = 50_000
TOTAL = 228_000


def money(n):
    return f"US${n:,.0f}"


def r(p, text, font="Georgia", size=10, bold=False, color=INK, italic=False):
    run = p.add_run(text)
    run.font.name = font
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    return run


def shade(cell, hexc):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexc)
    tcpr.append(shd)


def valign(cell, val="center"):
    tcpr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:vAlign")
    node.set(qn("w:val"), val)
    tcpr.append(node)


def set_cell_margins(cell, top=30, bottom=30, left=70, right=70):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def no_border_table(table):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblpr.append(borders)


def grid_borders(table, color="C5D5CE", sz="4"):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def set_table_widths(table, widths):
    """widths in inches. Fixed layout so Word and LibreOffice agree."""
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    total = int(sum(widths) * 1440)
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tblpr.addnext(grid)
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w * 1440)))
        grid.append(gc)
    for row in table.rows:
        tr = row._tr
        trpr = tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
        for i, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(int(widths[i] * 1440)))
            tcw.set(qn("w:type"), "dxa")


def tighten(p, before=0, after=0, exact=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0 if exact is None else Pt(exact)


def body(doc, text, size=10, after=2, before=0, exact=13):
    p = doc.add_paragraph()
    tighten(p, before, after, exact)
    r(p, text, "Georgia", size, False, INK)
    return p


def heading(doc, text, size=12, before=7, after=2):
    p = doc.add_paragraph()
    tighten(p, before, after, size + 2)
    p.paragraph_format.keep_with_next = True
    r(p, text, "Arial", size, True, GREEN)
    return p


def fill_cell(cell, text, size=8, bold=False, color=INK, fill=None, font="Arial", align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    tighten(p, 0, 0, size + 2)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r(p, text, font, size, bold, color)
    if fill:
        shade(cell, fill)
    valign(cell, "center")
    set_cell_margins(cell)


def add_bar(cell, fraction, max_width_in):
    """Green bar on a pale track. Width follows the stated dollar amount."""
    cell.text = ""
    shade(cell, "E3ECE8")
    valign(cell, "center")
    set_cell_margins(cell, top=36, bottom=36, left=50, right=50)
    p = cell.paragraphs[0]
    tighten(p, 0, 0, 2)
    width = max(0.12, max_width_in * fraction)
    inner = cell.add_table(1, 1)
    set_table_widths(inner, [width])
    no_border_table(inner)
    ic = inner.cell(0, 0)
    ic.text = ""
    shade(ic, "335D4F")
    ip = ic.paragraphs[0]
    tighten(ip, 0, 0, 8)
    r(ip, " ", "Arial", 4, False, WHITE)
    set_cell_margins(ic, top=0, bottom=0, left=0, right=0)
    # Word requires a paragraph after a nested table. Keep it from adding a line.
    extra = cell.paragraphs[-1]
    tighten(extra, 0, 0, 2)
    extra.paragraph_format.space_before = Pt(0)
    extra.paragraph_format.space_after = Pt(0)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.top_margin = Inches(0.40)
    section.bottom_margin = Inches(0.42)
    section.header_distance = Inches(0.16)
    section.footer_distance = Inches(0.16)

    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Georgia")

    content_w = 7.4

    logo_p = doc.add_paragraph()
    tighten(logo_p, 0, 1)
    logo_p.add_run().add_picture(str(LOGO), width=Inches(2.85))

    title = doc.add_paragraph()
    tighten(title, 1, 0, 16)
    title.paragraph_format.keep_with_next = True
    r(
        title,
        "REIN Ghana Hub \u00d7 Academic City: what ACity brings and what it costs",
        "Arial",
        13.5,
        True,
        GREEN,
    )

    sub = doc.add_paragraph()
    tighten(sub, 0, 2, 12)
    r(
        sub,
        "Internal prep for Gideon Blaauw (CleantechHUB) with Laurence Kalinsky (PVBLIC), "
        "before the meeting with Pav, Academic City University (ACity), Accra. "
        "20 October 2026.",
        "Arial",
        8.5,
        False,
        ORANGE,
    )

    heading(doc, "1. What ACity brings, in kind", before=4, after=1)

    body(
        doc,
        "TEC (Technology & Entrepreneurship Center), led by Makafui Awuku, runs the "
        "US State Department CCE Climate Tech Accelerator.",
    )
    body(
        doc,
        "In Q3 2025, TEC trained 140 youth in 73 workshops. Participation was 48% female.",
    )
    body(
        doc,
        "Africa's first Sustainability Laboratory was inaugurated on 9 February 2023. "
        "It is a CCE initiative, and home of the Telesto Climate Tech Accelerator "
        "(17 entrepreneurs in the 2023 cohort).",
    )
    body(
        doc,
        "About 500 students, and 25+ nationalities. AI and Robotics degrees fit "
        "PVBLIC ALI (AI Leapfrog Initiative) Logic Centers. International industry "
        "mentors 200+ students and educators a year. Northeastern University is a "
        "research link. Alumni are at Apple UK and Zipline.",
    )
    body(
        doc,
        "The US State Department CCE stays the anchor. REIN adds a non-US capital "
        "channel: FOSD (Family Offices for Sustainable Development) family offices "
        "and European corporates. That channel points toward the ACU Foundation's "
        "US$10 million goal over three years.",
        after=3,
    )

    stage_w = [1.35, 2.85, 3.20]
    stages = [
        ("REIN stage", "Scale", "ACity asset", True),
        ("Ideation", "50 participants, 4 weeks", "TEC workshops and students", False),
        ("Incubation", "25 teams, 10 weeks", "Sustainability Lab and Telesto", False),
        ("Acceleration", "10 teams, 12 weeks", "CCE Climate Tech Accelerator", False),
        (
            "Growth",
            "4 teams, 12 weeks, investor connections",
            "FOSD and PVBLIC / UNGA",
            False,
        ),
        ("Hub operations", "", "Campus, TEC staff", False),
    ]
    stage = doc.add_table(rows=len(stages), cols=3)
    set_table_widths(stage, stage_w)
    grid_borders(stage)
    for i, (a, b, c, hdr) in enumerate(stages):
        vals = (a, b, c)
        for j, val in enumerate(vals):
            if hdr:
                fill_cell(stage.cell(i, j), val, size=8, bold=True, color=WHITE, fill="335D4F")
            else:
                fill_cell(stage.cell(i, j), val, size=8, bold=(j == 0), color=INK)

    heading(doc, "2. Honest investment", before=7, after=1)
    body(doc, "These assets reduce cost and risk. They do not replace cash.")
    body(
        doc,
        "External cash follows a three-year sliding scale. Year 1 is the full funnel, "
        "plus setup and integration.",
        after=3,
    )

    bar_max = 3.15
    cash_w = [1.15, 1.45, 4.80]
    cash_rows = [
        ("Year", "External cash", None, True, False),
        ("Year 1", money(YEAR1), YEAR1 / YEAR1, False, False),
        ("Year 2", money(YEAR2), YEAR2 / YEAR1, False, False),
        ("Year 3", money(YEAR3), YEAR3 / YEAR1, False, False),
        ("3-year total", money(TOTAL), None, False, True),
    ]
    cash = doc.add_table(rows=len(cash_rows), cols=3)
    set_table_widths(cash, cash_w)
    grid_borders(cash, color="C5D5CE")
    for i, (label, amount, frac, hdr, total) in enumerate(cash_rows):
        if hdr:
            fill_cell(cash.cell(i, 0), label, size=8, bold=True, color=WHITE, fill="335D4F")
            fill_cell(cash.cell(i, 1), amount, size=8, bold=True, color=WHITE, fill="335D4F")
            fill_cell(cash.cell(i, 2), "Level, Year 1 = full bar", size=8, bold=True, color=WHITE, fill="335D4F")
        elif total:
            fill_cell(cash.cell(i, 0), label, size=8, bold=True, color=WHITE, fill="FF6633")
            fill_cell(cash.cell(i, 1), amount, size=8, bold=True, color=WHITE, fill="FF6633")
            fill_cell(cash.cell(i, 2), "Sum of the three years", size=8, bold=True, color=WHITE, fill="FF6633")
        else:
            fill_cell(cash.cell(i, 0), label, size=8, bold=True, color=INK)
            fill_cell(cash.cell(i, 1), amount, size=8, bold=True, color=INK)
            add_bar(cash.cell(i, 2), frac, bar_max)

    body(
        doc,
        "Year 1 split: programs US$75,000, comms and hub maintenance US$8,000, "
        "and operations and integration US$20,000.",
        before=3,
    )
    body(
        doc,
        "External need falls about 50% by Year 3, as ACity takes on delivery "
        "and local sponsors come in.",
    )

    heading(doc, "3. To agree with Laurence before Pav", before=6, after=1)
    agree = [
        (
            "a. ",
            "Year 1 cost split across PVBLIC, ACity, and FOSD. Unconfirmed.",
            False,
        ),
        (
            "b. ",
            "Can Pav put a value on the in-kind contribution? ",
            True,
        ),
        (
            "c. ",
            "What reduces in Years 2 and 3? The breakdown is not yet set. ",
            True,
        ),
    ]
    for label, text, confirm in agree:
        p = doc.add_paragraph()
        tighten(p, 1, 0, 13)
        r(p, label, "Arial", 10, True, GREEN)
        r(p, text, "Georgia", 10, False, INK)
        if confirm:
            r(p, "TO CONFIRM", "Arial", 10, True, ORANGE)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    tighten(fp, 2, 0, 11)
    ppr = fp._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "6")
    top.set(qn("w:space"), "1")
    top.set(qn("w:color"), "335D4F")
    pbdr.append(top)
    ppr.append(pbdr)
    r(
        fp,
        "Internal prep, not for circulation. Source: REIN \u00d7 Academic City proposal, Apr 2026.",
        "Arial",
        8,
        False,
        GREEN,
    )

    doc.core_properties.title = (
        "REIN Ghana Hub \u00d7 Academic City: what ACity brings and what it costs"
    )
    doc.core_properties.subject = "Internal prep, not for circulation"
    doc.core_properties.category = "REIN Hubs internal"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
